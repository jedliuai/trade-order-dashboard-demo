"""Real SQLite/HTTP regression tests; no live business-service dependencies."""
import copy
import http.client
import importlib
import io
import json
from pathlib import Path
import tempfile
import threading
import unittest
from datetime import date
from unittest.mock import patch


def fixture():
    owner = 'demo-sales-1'
    def rows(*items):
        return [dict(owner_id=owner, **item) for item in items]
    return {
        'customers': rows({'id': 'cu', 'name': '虚构客户', 'country': '示例国', 'default_currency': 'USD'},
                          {'id': 'cu2', 'name': '另一虚构客户', 'default_currency': 'USD'}),
        'contracts': rows({'id': 'c', 'contract_no': 'DEMO-C', 'customer_id': 'cu', 'currency': 'USD', 'incoterm': 'FOB', 'export_type': '自营', 'contract_date': '2026-08-02'},
                          {'id': 'c2', 'contract_no': 'DEMO-C2', 'customer_id': 'cu', 'currency': 'USD', 'incoterm': 'FOB', 'export_type': '自营', 'contract_date': '2026-07-02'},
                          {'id': 'c3', 'contract_no': 'DEMO-C3', 'customer_id': 'cu2', 'currency': 'USD', 'incoterm': 'FOB', 'export_type': '自营', 'contract_date': '2026-08-02'}),
        'contact_sheets': rows({'id': 'cs', 'contract_id': 'c', 'contact_sheet_no': 'DEMO-CS', 'material_no': 'DEMO-M', 'product_name': '虚构产品', 'specification': '10 ml', 'quantity': 100, 'unit_price': 10, 'unit': '瓶', 'business_type': '制剂'},
                              {'id': 'cs2', 'contract_id': 'c2', 'contact_sheet_no': 'DEMO-CS2', 'material_no': 'DEMO-M', 'product_name': '虚构产品', 'quantity': 50, 'unit_price': 10, 'unit': '瓶', 'business_type': '原料药'},
                              {'id': 'cs3', 'contract_id': 'c3', 'contact_sheet_no': 'DEMO-CS3', 'material_no': 'DEMO-M', 'product_name': '虚构产品', 'quantity': 50, 'unit_price': 10, 'unit': '瓶', 'business_type': '原料药'}),
        'batches': rows({'id': 'b', 'contact_sheet_id': 'cs', 'batch_no': 'DEMO-B', 'batch_quantity': 100, 'warehouse_date': '2026-08-03'}),
        'exchange_rates': [{'id': 'r', 'owner_id': 'demo-shared', 'effective_month': '2026-08', 'currency_pair': 'USD/CNY', 'rate': 7}],
        'product_costs': [{'id': 'cost', 'owner_id': 'demo-shared', 'material_no': 'DEMO-M', 'cost_month': '2026-08', 'unit_cost_no_tax': 20}],
    }


class LocalApiTest(unittest.TestCase):
    def setUp(self):
        try:
            self.store_module = importlib.import_module('local_api.store')
        except ModuleNotFoundError:
            self.fail('Local SQLite implementation is not present yet (intentional RED)')
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.seed_path = self.root / 'seed.json'
        self.seed_path.write_text(json.dumps(fixture(), ensure_ascii=False), encoding='utf-8')
        self.path = self.root / 'demo.sqlite3'
        self.store = self.store_module.Store(self.path, seed_path=self.seed_path)

    def rpc(self, name, **args):
        return self.store.rpc(name, args)

    def shipment(self, quantity=40, status='已发货', number='DEMO-S'):
        return self.rpc('create_shipment_with_items', shipment_data={
            'contract_id': 'c', 'shipment_no': number, 'shipment_date': '2026-08-04', 'status': status},
            item_rows=[{'contact_sheet_id': 'cs', 'batch_id': 'b', 'shipped_quantity': quantity, 'unit_price': 10}])

    def test_seed_volume_history_relations_and_role_coverage(self):
        build_seed = importlib.import_module('local_api.seed').build_seed
        data = build_seed(date(2026, 8, 27))
        self.assertGreaterEqual(len(data['customers']), 60)
        self.assertGreaterEqual(len(data['contracts']), 1000)
        self.assertGreaterEqual(len(data['contact_sheets']), 2000)
        self.assertLessEqual(min(r['contract_date'] for r in data['contracts']), '2024-08-01')
        self.assertTrue(any(r['contract_date'].startswith('2026-08') for r in data['contracts']))
        owners = {r['owner_id'] for r in data['contracts']}
        self.assertEqual(owners, {'demo-sales-1', 'demo-sales-2', 'demo-sales-3'})
        contracts = {r['id']: r for r in data['contracts']}
        customers = {r['id']: r for r in data['customers']}
        for row in data['contact_sheets']:
            parent = contracts[row['contract_id']]
            self.assertEqual(row['owner_id'], parent['owner_id'])
            self.assertEqual(parent['owner_id'], customers[parent['customer_id']]['owner_id'])
        self.assertEqual(data, build_seed(date(2026, 8, 27)))

    def test_customer_update_survives_reopen_and_projection(self):
        self.store.update('customers', {'id': 'eq.cu'}, {'name': '已修改虚构客户'})
        reopened = self.store_module.Store(self.path)
        self.assertEqual(reopened.query('customers', {'id': 'eq.cu', 'select': 'id,name'}),
                         [{'id': 'cu', 'name': '已修改虚构客户'}])

    def test_unknown_tables_actions_queries_and_empty_mutation_rejected(self):
        for table in ['auth.users', '../.env.local', 'sqlite_master']:
            with self.assertRaises(self.store_module.ApiError):
                self.store.query(table)
        with self.assertRaises(self.store_module.ApiError):
            self.rpc('not_implemented')
        with self.assertRaises(self.store_module.ApiError):
            self.store.query('customers', {'id': 'regex..*'})
        with self.assertRaises(self.store_module.ApiError):
            self.store.delete('customers', {})

    def test_roles_owner_spoofing_and_cross_owner_relationships(self):
        other = self.store.insert('customers', {'name': '乙客户'}, user='demo-sales-2')[0]
        self.assertEqual(self.store.query('customers', {'id': 'eq.' + other['id']}), [])
        self.assertEqual(len(self.store.query('customers', user='demo-manager')), 3)
        for user in ['demo-manager', 'demo-owner', 'unknown']:
            with self.assertRaises(self.store_module.ApiError):
                self.store.insert('customers', {'name': '禁止'}, user=user)
        with self.assertRaises(self.store_module.ApiError):
            self.store.insert('customers', {'owner_id': 'demo-sales-2', 'name': '伪造'})
        with self.assertRaises(self.store_module.ApiError):
            self.store.update('contracts', {'id': 'eq.c'}, {'customer_id': other['id']})

    def test_bulk_insert_rollback_and_foreign_key_delete_restriction(self):
        before = self.store.snapshot()
        with self.assertRaises(self.store_module.ApiError):
            self.store.insert('payments', [{'contract_id': 'c', 'amount': 20, 'payment_date': '2026-08-03', 'payment_type': '预付款'},
                                           {'contract_id': 'missing', 'amount': 1, 'payment_type': '其他'}])
        self.assertEqual(before, self.store.snapshot())
        with self.assertRaises(self.store_module.ApiError):
            self.store.delete('customers', {'id': 'eq.cu'})

    def test_batch_totals_tracking_and_replacement_rollback(self):
        with self.assertRaises(self.store_module.ApiError):
            self.rpc('replace_contact_sheet_batches', target_contact_sheet_id='cs', batch_rows=[{'batch_no': 'NEW', 'batch_quantity': 99}])
        self.assertEqual(self.store.query('batches')[0]['id'], 'b')
        self.assertEqual(self.rpc('update_batch_tracking', target_batch_id='b', target_warehouse_date='2026-08-05', target_release_date='2026-08-06', target_notes='已放行'), 1)
        self.assertEqual(self.store.query('batches')[0]['release_date'], '2026-08-06')
        with self.assertRaises(self.store_module.ApiError):
            self.store.update('contact_sheets', {'id': 'eq.cs'}, {'quantity': 101})

    def test_preparing_reserves_capacity_but_not_shipped_or_profit(self):
        self.shipment(60, '准备中')
        self.assertEqual(self.store.query('contact_sheets', {'id': 'eq.cs'})[0]['shipped_quantity'], 0)
        before = self.store.snapshot()
        with self.assertRaises(self.store_module.ApiError):
            self.shipment(41, number='OVER')
        self.assertEqual(before, self.store.snapshot())
        self.assertEqual(self.store.query('v_shipment_profit'), [])

    def test_atomic_shipment_update_and_cancel_releases_capacity(self):
        sid = self.shipment()
        self.assertEqual(self.rpc('update_shipment_with_items', target_shipment_id=sid,
            shipment_data={'shipment_no': 'REVISED', 'status': '取消', 'shipment_date': '2026-08-04'},
            target_payment_status='需人工确认', item_rows=[{'contact_sheet_id': 'cs', 'batch_id': 'b', 'shipped_quantity': 40, 'unit_price': 12}]), 1)
        self.shipment(100, number='FULL')
        before = self.store.snapshot()
        with self.assertRaises(self.store_module.ApiError):
            self.store.update('shipments', {'id': 'eq.' + sid}, {'status': '已发货'})
        self.assertEqual(before, self.store.snapshot())

    def test_group_cross_customer_rollback_and_actual_dispatch(self):
        items = [{'contact_sheet_id': 'cs', 'batch_id': 'b', 'shipped_quantity': 20, 'unit_price': 10},
                 {'contact_sheet_id': 'cs2', 'shipped_quantity': 10, 'unit_price': 10}]
        gid = self.rpc('create_shipment_group_with_items', group_data={'customer_id': 'cu', 'shipment_no': 'GROUP', 'status': '准备中'}, item_rows=items)
        self.assertEqual(len(self.store.query('shipments')), 2)
        before = self.store.snapshot()
        with self.assertRaises(self.store_module.ApiError):
            self.rpc('confirm_shipment_group_dispatched', target_group_id=gid, actual_shipment_date='2026-08-07', item_rows=[dict(items[1], contact_sheet_id='cs3')])
        self.assertEqual(before, self.store.snapshot())
        self.assertEqual(self.rpc('confirm_shipment_group_dispatched', target_group_id=gid, actual_shipment_date='2026-08-07', item_rows=items), gid)
        report = self.rpc('get_operating_metrics', p_start_date='2026-08-01', p_end_date='2026-08-31')
        self.assertEqual(report['shipments']['count'], 1)
        self.assertEqual(report['shipments']['amount_rmb'], 2100)

    def test_receipt_allocations_deposit_updates_history_and_no_double_count(self):
        rid = self.rpc('create_payment_receipt_with_allocations', receipt_data={'customer_id': 'cu', 'receipt_no': 'R1', 'payment_date': '2026-08-05', 'total_amount': 600, 'currency': 'USD', 'payment_type': '预付款'},
                       allocation_rows=[{'contract_id': 'c', 'amount': 400, 'payment_type': '尾款'}])
        report = self.rpc('get_operating_metrics', p_start_date='2026-08-01', p_end_date='2026-08-31')
        self.assertEqual(report['payments'], {'count': 1, 'original_usd': 600, 'original_rmb': 0, 'amount_rmb': 4200})
        self.assertEqual(report['customer_balance']['we_owe_customer_rmb'], 4200)
        before = self.store.snapshot()
        with self.assertRaises(self.store_module.ApiError):
            self.rpc('update_payment_receipt_with_allocations', target_receipt_id=rid, receipt_data={'total_amount': 600}, allocation_rows=[{'contract_id': 'c', 'amount': 601, 'payment_type': '尾款'}])
        self.assertEqual(before, self.store.snapshot())
        self.rpc('update_payment_receipt_with_allocations', target_receipt_id=rid, receipt_data={'total_amount': 700, 'change_reason': '虚构修改'}, allocation_rows=[{'contract_id': 'c', 'amount': 500, 'payment_type': '尾款'}])
        self.assertTrue(self.store.query('payment_receipt_change_history'))
        pid = self.store.query('payments')[0]['id']
        self.assertEqual(self.rpc('return_payment_allocation_to_deposit', target_payment_id=pid), rid)
        self.assertEqual(self.store.query('payments'), [])
        self.assertEqual(self.store.query('payment_receipts')[0]['total_amount'], 700)

    def test_profit_requires_invoice_exact_cost_and_uses_original_tax_rule(self):
        sid = self.shipment()
        self.assertEqual(self.store.query('v_shipment_profit'), [])
        self.rpc('mark_physical_shipment_invoiced', target_shipment_id=sid, invoice_data={'invoice_date': '2026-08-08'})
        profit = self.store.query('v_shipment_profit')[0]
        self.assertEqual(profit['profit'], 2000)
        self.assertFalse(profit['is_estimated_profit'])
        self.assertEqual(self.rpc('get_operating_metrics', p_start_date='2026-08-01', p_end_date='2026-08-31')['confirmed_profit_rmb'], 2000)
        with self.store.transaction() as tx:
            tx.rows['product_costs'][0]['cost_month'] = '2026-07'
        self.assertTrue(self.store.query('v_shipment_profit')[0]['is_estimated_profit'])
        self.assertEqual(self.rpc('get_operating_metrics', p_start_date='2026-08-01', p_end_date='2026-08-31')['confirmed_profit_rmb'], 0)
        with self.store.transaction() as tx:
            tx.rows['exchange_rates'].clear()
        self.assertIsNone(self.store.query('v_shipment_profit')[0]['profit'])

    def test_cif_freight_and_invoice_edit_rollback(self):
        self.store.update('contracts', {'id': 'eq.c'}, {'incoterm': 'CIF'})
        sid = self.shipment()
        with self.assertRaises(self.store_module.ApiError):
            self.rpc('mark_physical_shipment_invoiced', target_shipment_id=sid, invoice_data={'invoice_date': '2026-08-08'})
        iid = self.rpc('mark_physical_shipment_invoiced', target_shipment_id=sid, invoice_data={'invoice_date': '2026-08-08', 'freight_insurance_amount': 40})
        profit = self.store.query('v_shipment_profit')[0]
        self.assertEqual(profit['sales_amount'], 400)
        self.assertEqual(profit['profit'], 1720)
        before = self.store.snapshot()
        with self.assertRaises(self.store_module.ApiError):
            self.rpc('update_physical_shipment_invoice', target_invoice_id=iid, invoice_data={'invoice_date': '2026-08-09', 'freight_insurance_amount': 400})
        self.assertEqual(before, self.store.snapshot())

    def test_invoice_completion_and_archive_reversible(self):
        sid = self.shipment(100)
        self.store.insert('payments', {'contract_id': 'c', 'payment_date': '2026-08-05', 'amount': 1000, 'payment_type': '尾款'})
        iid = self.rpc('create_invoice_with_shipments', invoice_data={'amount': 1000, 'status': '已申请'}, shipment_rows=[{'shipment_id': sid, 'allocated_amount': 1000}])
        self.assertFalse(self.store.query('contracts', {'id': 'eq.c'})[0]['archived'])
        self.store.update('invoices', {'id': 'eq.' + iid}, {'invoice_date': '2026-08-10'})
        self.assertTrue(self.store.query('contracts', {'id': 'eq.c'})[0]['archived'])
        self.store.delete('invoices', {'id': 'eq.' + iid})
        self.assertFalse(self.store.query('contracts', {'id': 'eq.c'})[0]['archived'])

    def test_partial_dated_invoice_does_not_archive_fully_paid_contract(self):
        sid = self.shipment(100)
        self.store.insert('payments', {'contract_id': 'c', 'payment_date': '2026-08-05',
                          'amount': 1000, 'payment_type': '尾款'})
        self.rpc('create_invoice_with_shipments', invoice_data={
            'amount': 1, 'invoice_date': '2026-08-10'},
            shipment_rows=[{'shipment_id': sid, 'allocated_amount': 1}])
        self.assertFalse(self.store.query('contracts', {'id': 'eq.c'})[0]['archived'])

    def test_master_data_relations_versions_packaging_and_preferences(self):
        product = self.store.insert('mdc_products', {'chinese_name': '虚构品', 'product_type': 'finished_product'})[0]
        variant = self.store.insert('mdc_product_variants', {'product_id': product['id'], 'specification': '10ml', 'status': 'active'})[0]
        self.store.update('contact_sheets', {'id': 'eq.cs'}, {'product_variant_id': variant['id']})
        revised = self.rpc('revise_mdc_product_variant', target_variant_id=variant['id'], replacement_specification='20ml')
        self.assertTrue(revised['preserved_historical_variant'])
        self.assertNotEqual(revised['variant_id'], variant['id'])
        self.assertEqual(self.store.query('contact_sheets', {'id': 'eq.cs'})[0]['product_variant_id'], variant['id'])
        nested = self.store.query('mdc_products', {'select': '*,aliases:mdc_product_aliases(id,alias,alias_kind,created_at),variants:mdc_product_variants(id,specification,status)'})[0]
        self.assertEqual(len(nested['variants']), 2)
        pid = self.rpc('create_packaging_profile_with_version', profile_data={'product_name': '虚构品', 'product_variant_id': revised['variant_id'], 'is_active': True}, version_data={'quantity_per_carton': 100, 'quantity_unit': '瓶/箱', 'packaging_description': '虚构纸箱'})
        vid = self.rpc('create_packaging_version', target_profile_id=pid, profile_data={}, version_data={'quantity_per_carton': 120, 'quantity_unit': '瓶/箱'})
        current = self.store.query('packaging_current_profiles')[0]
        self.assertEqual(current['version_id'], vid)
        self.assertEqual(current['version_no'], 2)
        self.store.insert('user_ui_preferences', {'customer_display_order': ['cu2', 'cu']}, params={'on_conflict': 'owner_id'})
        self.store.insert('user_ui_preferences', {'customer_display_order': ['cu']}, params={'on_conflict': 'owner_id'})
        self.assertEqual(self.store.query('user_ui_preferences', {'select': 'customer_display_order'}), [{'customer_display_order': ['cu']}])

    def test_alert_refresh_preserves_snooze_and_history(self):
        sid = self.shipment()
        self.rpc('refresh_alerts')
        alerts = self.store.query('alerts', {'related_id': 'eq.' + sid})
        self.assertTrue(alerts)
        self.store.update('alerts', {'id': 'eq.' + alerts[0]['id']}, {'status': '稍后提醒', 'snoozed_until': '2099-01-01'})
        self.rpc('refresh_alerts')
        self.assertEqual(self.store.query('alerts', {'id': 'eq.' + alerts[0]['id']})[0]['status'], '稍后提醒')
        self.store.update('contact_sheets', {'id': 'eq.cs'}, {'notes': '虚构变更'})
        self.assertTrue(self.store.query('contact_sheet_change_history'))

    def test_exports_reopen_all_ten_types_and_apply_filters(self):
        exports = importlib.import_module('local_api.exports')
        from openpyxl import load_workbook
        from docx import Document
        gid = self.rpc('create_shipment_group_with_items', group_data={'customer_id': 'cu', 'shipment_no': 'GROUP', 'status': '已发货', 'shipment_date': '2026-08-04'}, item_rows=[{'contact_sheet_id': 'cs', 'batch_id': 'b', 'shipped_quantity': 10, 'unit_price': 10}])
        sid = self.store.query('shipments')[0]['id']
        self.rpc('mark_physical_shipment_invoiced', target_shipment_id=sid, invoice_data={'invoice_date': '2026-08-05'})
        self.store.insert('payments', {'contract_id': 'c', 'amount': 50, 'payment_date': '2026-08-05', 'payment_type': '尾款'})
        commands = ['export-production', 'export-monthly', 'export-original-contracts', 'export-helper', 'export-payments', 'export-customer-monthly-sales', 'export-rmb-invoice', 'export-shipment-plan', 'export-receipt-confirmation', 'export-shipment-details']
        for command in commands:
            with self.subTest(command=command):
                payload = {'customer_id': 'cu', 'contract_id': 'c', 'date_from': '2026-08-01', 'date_to': '2026-08-31', 'as_of_date': '2026-08-27', 'shipment_group_id': gid, 'shipment_id': sid, 'contact_sheet_ids': ['cs']}
                name, mime, content = exports.export_file(self.store, command, payload)
                if name.endswith('.docx'):
                    doc = Document(io.BytesIO(content))
                    text = '\n'.join(p.text for p in doc.paragraphs) + '\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
                    self.assertIn('虚构客户', text)
                    self.assertIn('DEMO-CS', text)
                else:
                    workbook = load_workbook(io.BytesIO(content), data_only=False)
                    text = str(list(workbook.active.values))
                    self.assertTrue(workbook.active.max_column > 2)
                    self.assertNotIn('DEMO-C2', text)
                    if command != 'export-rmb-invoice':
                        self.assertGreater(workbook.active.max_row, 1)
                    workbook.close()
        _, _, content = exports.export_file(self.store, 'export-original-contracts', {'date_from': '2026-08-01', 'date_to': '2026-08-31'})
        workbook = load_workbook(io.BytesIO(content))
        self.assertIn('DEMO-C', str(list(workbook.active.values)))
        self.assertNotIn('DEMO-C2', str(list(workbook.active.values)))
        workbook.close()

    def test_telegram_preview_shares_metrics_and_never_sends_unconfigured(self):
        telegram = importlib.import_module('local_api.telegram')
        self.shipment()
        preview = telegram.preview(self.store, {'mode': 'current_month', 'recipient': 'manager'}, as_of=date(2026, 8, 27))
        self.assertEqual(preview['report']['shipments']['amount_rmb'], 2800)
        self.assertIn('演示', preview['text'])
        fiscal = telegram.preview(self.store, {'mode': 'fiscal_year', 'recipient': 'owner'}, as_of=date(2026, 8, 27))
        self.assertEqual(fiscal['report']['period']['start_date'], '2025-12-01')
        with patch('urllib.request.urlopen', side_effect=AssertionError('network forbidden')):
            with self.assertRaises(self.store_module.ApiError):
                telegram.send(self.store, {'mode': 'current_month', 'recipient': 'manager', 'confirmed': True}, env_path=self.root / '.env.local')
            with self.assertRaises(self.store_module.ApiError):
                telegram.send(self.store, {'mode': 'current_month', 'recipient': 'manager'}, env_path=self.root / '.env.local')

    def test_http_json_security_roles_download_and_reset_backup(self):
        server = importlib.import_module('local_api.server')
        httpd = server.make_server(self.store, port=0, env_path=self.root / '.env.local')
        thread = threading.Thread(target=httpd.serve_forever, daemon=True)
        thread.start()
        self.addCleanup(httpd.server_close)
        self.addCleanup(httpd.shutdown)
        def request(method, path, body=None, headers=None):
            connection = http.client.HTTPConnection('127.0.0.1', httpd.server_port, timeout=5)
            try:
                data = json.dumps(body) if body is not None else None
                connection.request(method, path, data, {'Content-Type': 'application/json', **(headers or {})})
                response = connection.getresponse()
                return response.status, dict(response.getheaders()), response.read()
            finally:
                connection.close()
        status, _, body = request('GET', '/api/health')
        self.assertEqual(status, 200)
        self.assertEqual(json.loads(body)['status'], 'ok')
        self.assertEqual(request('POST', '/api/data/customers', {'name': 'bad'}, {'X-Demo-User': 'demo-manager'})[0], 403)
        self.assertEqual(request('GET', '/api/data/customers', headers={'Host': 'evil.example'})[0], 403)
        self.assertEqual(request('POST', '/api/data/customers', {'name': 'bad'}, {'Origin': 'https://evil.example'})[0], 403)
        self.assertEqual(request('GET', '/api/data/no_table')[0], 404)
        self.assertEqual(request('GET', '/api/data/rpc/refresh_alerts')[0], 405)
        status, headers, body = request('GET', '/api/data/customers?select=id,name&id=eq.cu')
        self.assertEqual(json.loads(body), [{'id': 'cu', 'name': '虚构客户'}])
        self.assertEqual(headers['Cache-Control'], 'no-store')
        self.assertEqual(request('POST', '/api/telegram/send', {'mode': 'current_month', 'recipient': 'manager', 'confirmed': True})[0], 503)
        status, headers, body = request('POST', '/api/export/export-original-contracts', {})
        self.assertEqual(status, 200)
        self.assertTrue(body.startswith(b'PK'))
        self.assertIn('attachment', headers['Content-Disposition'])
        self.store.update('customers', {'id': 'eq.cu'}, {'notes': 'backup me'})
        self.assertEqual(request('POST', '/api/reset', {'confirm': 'NO'})[0], 400)
        status, _, body = request('POST', '/api/reset', {'confirm': 'RESET_DEMO'})
        self.assertEqual(status, 200)
        backup = Path(json.loads(body)['backup_path'])
        snapshot = json.loads(backup.read_text(encoding='utf-8'))
        self.assertTrue(any(r.get('notes') == 'backup me' for r in snapshot['customers']))
        self.assertNotIn('backup me', str(self.store.query('customers')))
        status, _, body = request('GET', '/api/backup')
        self.assertEqual(status, 200)
        self.assertIn('contracts', json.loads(body))


if __name__ == '__main__':
    unittest.main()
