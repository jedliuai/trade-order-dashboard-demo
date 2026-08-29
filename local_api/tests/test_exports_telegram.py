"""Behaviour tests for locally generated exports and Telegram reporting."""
from datetime import date
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from local_api.store import ApiError, Store
from local_api.tests.test_local_api import fixture


class ExportAndTelegramTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        root = Path(self.temp.name)
        seed = root / 'seed.json'
        seed.write_text(json.dumps(fixture(), ensure_ascii=False), encoding='utf-8')
        self.store = Store(root / 'demo.sqlite3', seed_path=seed)
        self.env_path = root / '.env.local'

    def add_sheet_and_batch(self, *, sheet_id='cs-extra', batch_id='b-extra', contract_id='c', quantity=20, price=10):
        with self.store.transaction() as tx:
            tx.add('contact_sheets', {
                'id': sheet_id, 'contract_id': contract_id, 'contact_sheet_no': sheet_id.upper(),
                'material_no': 'DEMO-M', 'product_name': '另一虚构产品', 'specification': '20 ml',
                'quantity': quantity, 'unit_price': price, 'unit': '瓶', 'business_type': '制剂',
            })
            tx.add('batches', {'id': batch_id, 'contact_sheet_id': sheet_id,
                               'batch_no': batch_id.upper(), 'batch_quantity': quantity,
                               'warehouse_date': '2026-08-03'})
        return sheet_id, batch_id

    def export_rows(self, command, filters):
        from local_api.exports import export_file
        from openpyxl import load_workbook

        _, _, content = export_file(self.store, command, filters)
        workbook = load_workbook(io.BytesIO(content), data_only=False)
        values = list(workbook.active.values)
        workbook.close()
        headers = list(values[3])
        return [dict(zip(headers, row)) for row in values[4:] if any(value is not None for value in row)]

    def test_contract_export_uses_real_snapshot_filters_and_escapes_formula_text(self):
        from local_api.exports import export_file
        from openpyxl import load_workbook

        self.store.update('customers', {'id': 'eq.cu'}, {'name': '=untrusted formula'})
        name, mime, content = export_file(
            self.store,
            'export-original-contracts',
            {'customer_id': 'cu', 'date_from': '2026-08-01', 'date_to': '2026-08-31'},
        )

        self.assertTrue(name.endswith('.xlsx'))
        self.assertEqual(mime, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        workbook = load_workbook(io.BytesIO(content), data_only=False)
        values = list(workbook.active.values)
        workbook.close()
        text = str(values)
        self.assertIn('DEMO-C', text)
        self.assertNotIn('DEMO-C2', text)
        self.assertNotIn('DEMO-C3', text)
        self.assertIn("'=untrusted formula", text)

    def test_unknown_export_command_is_rejected(self):
        from local_api.exports import export_file

        with self.assertRaises(ApiError):
            export_file(self.store, 'export-not-a-command', {})

    def test_payment_export_filters_on_payment_date_not_contract_date(self):
        from local_api.exports import export_file
        from openpyxl import load_workbook

        self.store.insert('payments', {
            'contract_id': 'c2', 'amount': 50, 'payment_date': '2026-08-05',
            'payment_type': '尾款', 'exchange_rate': 7,
        })
        _, _, content = export_file(self.store, 'export-payments', {
            'date_from': '2026-08-01', 'date_to': '2026-08-31',
        })
        workbook = load_workbook(io.BytesIO(content))
        text = str(list(workbook.active.values))
        workbook.close()
        self.assertIn('DEMO-C2', text)

    def test_shipment_scope_that_matches_nothing_does_not_fall_back_to_all_contracts(self):
        from local_api.exports import export_file
        from openpyxl import load_workbook

        _, _, content = export_file(self.store, 'export-original-contracts', {'shipment_id': 'missing'})
        workbook = load_workbook(io.BytesIO(content))
        text = str(list(workbook.active.values))
        workbook.close()
        self.assertNotIn('DEMO-C', text)

    def test_shipment_detail_profit_is_mapped_per_contact_sheet_item(self):
        second_sheet, second_batch = self.add_sheet_and_batch()
        shipment_id = self.store.rpc('create_shipment_with_items', {
            'shipment_data': {'contract_id': 'c', 'shipment_no': 'MULTI', 'shipment_date': '2026-08-04', 'status': '已发货'},
            'item_rows': [
                {'contact_sheet_id': 'cs', 'batch_id': 'b', 'shipped_quantity': 10, 'unit_price': 10},
                {'contact_sheet_id': second_sheet, 'batch_id': second_batch, 'shipped_quantity': 20, 'unit_price': 10},
            ],
        })
        self.store.rpc('mark_physical_shipment_invoiced', {
            'target_shipment_id': shipment_id, 'invoice_data': {'invoice_date': '2026-08-05'},
        })

        rows = self.export_rows('export-shipment-details', {'shipment_id': shipment_id})
        profits = {row['联系单号']: row['利润（人民币）'] for row in rows}
        self.assertEqual(profits, {'DEMO-CS': 500, 'CS-EXTRA': 1000})

    def test_customer_monthly_uses_receipt_parent_and_keeps_prepayment_nonnegative(self):
        shipment_id = self.store.rpc('create_shipment_with_items', {
            'shipment_data': {'contract_id': 'c', 'shipment_no': 'PREPAY', 'shipment_date': '2026-08-04', 'status': '已发货'},
            'item_rows': [{'contact_sheet_id': 'cs', 'batch_id': 'b', 'shipped_quantity': 10, 'unit_price': 10}],
        })
        self.store.rpc('create_payment_receipt_with_allocations', {
            'receipt_data': {'customer_id': 'cu', 'total_amount': 300, 'currency': 'USD', 'payment_date': '2026-08-05', 'exchange_rate': 7, 'payment_type': '预付款'},
            'allocation_rows': [],
        })

        rows = self.export_rows('export-customer-monthly-sales', {'as_of_date': '2026-08-27'})
        result = next(row for row in rows if row['客户'] == '虚构客户')
        self.assertEqual(result['币种'], 'USD')
        self.assertEqual(result['本月回款原币'], 300)
        self.assertEqual(result['当前应收原币'], 0)
        self.assertEqual(result['客户预存款原币'], 200)
        self.assertEqual(result['当前应收人民币'], 0)

    def test_rmb_invoice_uses_invoice_bridge_for_selected_group_child(self):
        with self.store.transaction() as tx:
            tx.add('customers', {'id': 'rmb-customer', 'name': '人民币虚构客户', 'default_currency': 'RMB'})
            for suffix in ('one', 'two'):
                contract_id, sheet_id, batch_id = f'rmb-{suffix}', f'rmb-sheet-{suffix}', f'rmb-batch-{suffix}'
                tx.add('contracts', {'id': contract_id, 'contract_no': contract_id.upper(), 'customer_id': 'rmb-customer',
                                     'currency': 'RMB', 'incoterm': None, 'export_type': '转口', 'contract_date': '2026-08-01'})
                tx.add('contact_sheets', {'id': sheet_id, 'contract_id': contract_id, 'contact_sheet_no': sheet_id.upper(),
                                          'material_no': 'RMB-M', 'product_name': '人民币产品', 'specification': '1 kg',
                                          'quantity': 10, 'unit_price': 10, 'unit': 'kg', 'business_type': '原料药'})
                tx.add('batches', {'id': batch_id, 'contact_sheet_id': sheet_id, 'batch_no': batch_id.upper(), 'batch_quantity': 10})
        group_id = self.store.rpc('create_shipment_group_with_items', {
            'group_data': {'customer_id': 'rmb-customer', 'shipment_no': 'RMB-G', 'status': '已发货', 'shipment_date': '2026-08-04'},
            'item_rows': [
                {'contact_sheet_id': 'rmb-sheet-one', 'batch_id': 'rmb-batch-one', 'shipped_quantity': 10, 'unit_price': 10},
                {'contact_sheet_id': 'rmb-sheet-two', 'batch_id': 'rmb-batch-two', 'shipped_quantity': 10, 'unit_price': 10},
            ],
        })
        children = sorted(self.store.query('shipments'), key=lambda row: row['id'])
        selected, other = children[-1], children[0]
        self.store.rpc('create_invoice_with_shipments', {
            'invoice_data': {'amount': 200, 'status': '已申请'},
            'shipment_rows': [{'shipment_id': other['id'], 'allocated_amount': 100}, {'shipment_id': selected['id'], 'allocated_amount': 100}],
        })

        rows = self.export_rows('export-rmb-invoice', {'shipment_id': selected['id'], 'shipment_group_id': group_id})
        self.assertEqual([row['联系单号'] for row in rows], [selected['contract_id'].replace('rmb-', 'rmb-sheet-').upper()])
        self.assertEqual(rows[0]['发票状态'], '已申请')
        self.assertEqual(rows[0]['分摊开票金额'], 100)

    def test_production_helper_and_receipt_exports_preserve_operational_fields(self):
        self.store.update('contact_sheets', {'id': 'eq.cs'}, {
            'qa_approval_date': '2026-08-01', 'production_date': '2026-08-02', 'expiry_date': '2028-08-01',
            'packaging': '10 瓶/箱', 'quality_standard': '演示质标', 'reference_contact_sheet': 'REF-DEMO',
        })
        self.store.update('batches', {'id': 'eq.b'}, {'production_date': '2026-08-02', 'expiry_date': '2028-08-01'})
        production = self.export_rows('export-production', {'contact_sheet_ids': ['cs']})
        helper = self.export_rows('export-helper', {'contract_id': 'c'})
        self.assertEqual(production[0]['批号汇总'], 'DEMO-B')
        self.assertEqual(production[0]['包装/包材'], '10 瓶/箱')
        self.assertEqual(helper[0]['批次数量说明'], 'DEMO-B：100 瓶')
        shipment_id = self.store.rpc('create_shipment_with_items', {
            'shipment_data': {'contract_id': 'c', 'shipment_no': 'UNIT', 'shipment_date': '2026-08-04', 'status': '已发货'},
            'item_rows': [{'contact_sheet_id': 'cs', 'batch_id': 'b', 'shipped_quantity': 10, 'unit_price': 10}],
        })
        from local_api.exports import export_file
        from docx import Document
        _, _, document = export_file(self.store, 'export-receipt-confirmation', {'shipment_id': shipment_id})
        text = '\n'.join(cell.text for table in Document(io.BytesIO(document)).tables for row in table.rows for cell in row.cells)
        self.assertIn('瓶', text)
        self.assertIn('2026-08-04', text)

    def test_telegram_configuration_is_boolean_only_and_send_uses_telegram_boundary(self):
        from local_api import telegram

        self.env_path.write_text(
            'TELEGRAM_BOT_TOKEN=secret-token\n'
            'TELEGRAM_MANAGER_CHAT_ID=manager-chat\n'
            'TELEGRAM_OWNER_CHAT_ID=owner-chat\n',
            encoding='utf-8',
        )
        self.assertEqual(
            telegram.status(self.env_path),
            {'configured': True, 'recipients': {'manager': True, 'owner': True}},
        )
        with patch('urllib.request.urlopen') as urlopen:
            urlopen.return_value.__enter__.return_value.read.return_value = b'{"ok": true}'
            result = telegram.send(
                self.store,
                {'mode': 'fiscal_year', 'recipient': 'manager', 'confirmed': True},
                env_path=self.env_path,
                as_of=date(2026, 8, 27),
            )
        request = urlopen.call_args.args[0]
        self.assertEqual(request.full_url, 'https://api.telegram.org/botsecret-token/sendMessage')
        self.assertEqual(result['recipient'], 'manager')
        self.assertEqual(result['report']['period']['start_date'], '2025-12-01')

    def test_telegram_requires_literal_confirmation_and_does_not_expose_configuration(self):
        from local_api import telegram

        self.env_path.write_text('TELEGRAM_BOT_TOKEN=secret-token\n', encoding='utf-8')
        with self.assertRaises(ApiError) as caught:
            telegram.send(
                self.store,
                {'mode': 'current_month', 'recipient': 'manager', 'confirmed': 1},
                env_path=self.env_path,
            )
        self.assertNotIn('secret-token', str(caught.exception))


if __name__ == '__main__':
    unittest.main()
