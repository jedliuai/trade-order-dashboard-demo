"""Safe, generic local report generation for the demonstration database."""
from collections import defaultdict
from datetime import date
import io
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from .store import ApiError, iso_date


XLSX_MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
DOCX_MIME = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
COMMANDS = {
    'export-production', 'export-monthly', 'export-original-contracts', 'export-helper',
    'export-payments', 'export-customer-monthly-sales', 'export-rmb-invoice',
    'export-shipment-plan', 'export-receipt-confirmation', 'export-shipment-details',
}


def _text(value):
    """Render a cell as inert text so spreadsheet formulas are never executed."""
    if value is None:
        return ''
    if isinstance(value, str) and value[:1] in ('=', '+', '-', '@'):
        return "'" + value
    return value


def _ids(filters, key):
    if key not in filters:
        return None
    value = filters[key]
    if value in (None, ''):
        return None
    if isinstance(value, str):
        return {value}
    if isinstance(value, list) and all(isinstance(item, str) for item in value):
        return set(value)
    raise ApiError(f'{key} 必须是字符串或字符串数组')


def _date_range(filters):
    start, end = filters.get('date_from', ''), filters.get('date_to', '')
    if start:
        iso_date(start, '起始日期')
    if end:
        iso_date(end, '截止日期')
    if start and end and start > end:
        raise ApiError('起始日期不能晚于截止日期')
    return start, end


def _in_range(value, start, end):
    return bool(value) and (not start or value >= start) and (not end or value <= end)


def _context(store, filters, user):
    if not isinstance(filters, dict):
        raise ApiError('导出筛选必须为 JSON 对象')
    rows = store.snapshot(user)
    return {
        'rows': rows,
        'customers': {row['id']: row for row in rows['customers']},
        'contracts': {row['id']: row for row in rows['contracts']},
        'sheets': {row['id']: row for row in rows['contact_sheets']},
        'batches': {row['id']: row for row in rows['batches']},
        'shipments': {row['id']: row for row in rows['shipments']},
        'groups': {row['id']: row for row in rows['shipment_groups']},
        'filters': filters,
        'user': user,
    }


def _match_contract(contract, context, date_field='contract_date', apply_dates=True):
    filters = context['filters']
    customer_ids, contract_ids = _ids(filters, 'customer_id'), _ids(filters, 'contract_id')
    shipment_ids, group_ids = _ids(filters, 'shipment_id'), _ids(filters, 'shipment_group_id')
    start, end = _date_range(filters)
    if customer_ids is not None and contract.get('customer_id') not in customer_ids:
        return False
    if contract_ids is not None and contract.get('id') not in contract_ids:
        return False
    if filters.get('contract_no') and contract.get('contract_no') != filters['contract_no']:
        return False
    if shipment_ids is not None or group_ids is not None:
        linked = any(
            shipment.get('contract_id') == contract.get('id')
            and (shipment_ids is None or shipment.get('id') in shipment_ids)
            and (group_ids is None or shipment.get('shipment_group_id') in group_ids)
            for shipment in context['rows']['shipments']
        )
        if not linked:
            return False
    requested_sheets = _ids(filters, 'contact_sheet_ids')
    if requested_sheets is not None and not any(
        sheet.get('contract_id') == contract.get('id') and sheet.get('id') in requested_sheets
        for sheet in context['rows']['contact_sheets']
    ):
        return False
    return not apply_dates or not (start or end) or _in_range(contract.get(date_field), start, end)


def _match_sheet(sheet, context):
    requested = _ids(context['filters'], 'contact_sheet_ids')
    return requested is None or sheet.get('id') in requested


def _match_shipment(shipment, context):
    filters = context['filters']
    shipment_ids, group_ids = _ids(filters, 'shipment_id'), _ids(filters, 'shipment_group_id')
    if shipment_ids is not None and shipment.get('id') not in shipment_ids:
        return False
    if group_ids is not None and shipment.get('shipment_group_id') not in group_ids:
        return False
    contract = context['contracts'].get(shipment.get('contract_id'))
    if not contract or not _match_contract(contract, context, apply_dates=False):
        return False
    start, end = _date_range(filters)
    return not (start or end) or _in_range(shipment.get('shipment_date'), start, end)


def _line_rows(context, *, shipments_only=False):
    rows = []
    items_by_shipment = defaultdict(list)
    for item in context['rows']['shipment_items']:
        items_by_shipment[item.get('shipment_id')].append(item)
    if shipments_only:
        sources = [shipment for shipment in context['rows']['shipments'] if _match_shipment(shipment, context)]
        for shipment in sources:
            contract = context['contracts'][shipment['contract_id']]
            customer = context['customers'].get(contract.get('customer_id'), {})
            for item in items_by_shipment[shipment['id']]:
                sheet = context['sheets'].get(item.get('contact_sheet_id'), {})
                if _match_sheet(sheet, context):
                    batch = context['batches'].get(item.get('batch_id'), {})
                    rows.append({
                        '客户': customer.get('name'), '合同号': contract.get('contract_no'),
                        '联系单号': sheet.get('contact_sheet_no'), '产品名称': sheet.get('product_name'),
                        '批号': batch.get('batch_no'), '发货单号': shipment.get('shipment_no'),
                        '发货日期': shipment.get('shipment_date'), '发货状态': shipment.get('status'),
                        '数量': item.get('shipped_quantity'), '单价': item.get('unit_price'),
                        '金额': item.get('amount'), '单位': sheet.get('unit'), '币种': shipment.get('currency'),
                        '_shipment_id': shipment.get('id'), '_contact_sheet_id': sheet.get('id'),
                        '_batch_id': batch.get('id'),
                    })
        return rows
    for sheet in context['rows']['contact_sheets']:
        contract = context['contracts'].get(sheet.get('contract_id'))
        if not contract or not _match_contract(contract, context) or not _match_sheet(sheet, context):
            continue
        customer = context['customers'].get(contract.get('customer_id'), {})
        rows.append({
            '客户': customer.get('name'), '国家/地区': customer.get('country'), '合同号': contract.get('contract_no'),
            '签署日期': contract.get('contract_date'), '贸易方式': contract.get('export_type'),
            '联系单号': sheet.get('contact_sheet_no'), '产品名称': sheet.get('product_name'),
            '规格': sheet.get('specification'), '物料号': sheet.get('material_no'), '数量': sheet.get('quantity'),
            '单位': sheet.get('unit'), '单价': sheet.get('unit_price'),
            '合同金额': (sheet.get('quantity') or 0) * (sheet.get('unit_price') or 0),
            '已发货数量': sheet.get('shipped_quantity'), '业务类型': sheet.get('business_type'),
        })
    return rows


def _batches_by_sheet(context):
    grouped = defaultdict(list)
    for batch in context['rows']['batches']:
        grouped[batch.get('contact_sheet_id')].append(batch)
    return grouped


def _joined(values):
    return '、'.join(str(value) for value in values if value not in (None, ''))


def _operational_rows(context, kind):
    batches_by_sheet = _batches_by_sheet(context)
    rows = []
    for sheet in context['rows']['contact_sheets']:
        contract = context['contracts'].get(sheet.get('contract_id'))
        if not contract or not _match_contract(contract, context) or not _match_sheet(sheet, context):
            continue
        customer = context['customers'].get(contract.get('customer_id'), {})
        batches = batches_by_sheet[sheet.get('id')]
        batch_text = _joined(batch.get('batch_no') for batch in batches)
        production_date = sheet.get('production_date') or _joined(batch.get('production_date') for batch in batches)
        expiry_date = sheet.get('expiry_date') or _joined(batch.get('expiry_date') for batch in batches)
        batch_note = _joined(
            f"{batch.get('batch_no')}：{float(batch.get('batch_quantity') or 0):g} {sheet.get('unit') or ''}".strip()
            for batch in batches
        )
        if kind == 'production':
            rows.append({
                '联系单号': sheet.get('contact_sheet_no'), '合同号': contract.get('contract_no'), '客户': customer.get('name'),
                '产品名称': sheet.get('product_name'), '规格': sheet.get('specification'), '批号汇总': batch_text,
                '生产/排产日期': production_date, '有效期': expiry_date, 'QA审批日期': sheet.get('qa_approval_date'),
                '实际入库日期': sheet.get('actual_warehousing_date') or sheet.get('warehouse_date'),
                '包装/包材': sheet.get('packaging'), '数量': sheet.get('quantity'), '单位': sheet.get('unit'),
                '已发货数量': sheet.get('shipped_quantity'),
                '交付说明': f"客户指定交货日期：{contract.get('agreed_delivery_date')}" if contract.get('agreed_delivery_date') else '',
            })
        else:
            rows.append({
                '联系单号': sheet.get('contact_sheet_no'), '产品名称': sheet.get('product_name'),
                '规格': sheet.get('specification'), '批号汇总': batch_text, '生产日期': production_date,
                '有效期': expiry_date, '数量': sheet.get('quantity'), '单位': sheet.get('unit'),
                '包装/包材': sheet.get('packaging'), '参考联系单': sheet.get('reference_contact_sheet'),
                '质量标准': sheet.get('quality_standard'), '批次数量说明': batch_note,
            })
    return rows


def _xlsx(title, subtitle, headers, rows):
    book = Workbook()
    sheet = book.active
    sheet.title = re.sub(r'[:\\/*?\[\]]', '_', title)[:31] or '演示报表'
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, len(headers)))
    heading = sheet.cell(1, 1, _text(title))
    heading.font = Font(size=16, bold=True, color='FFFFFF')
    heading.fill = PatternFill('solid', fgColor='1F4E78')
    heading.alignment = Alignment(horizontal='center')
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(1, len(headers)))
    sheet.cell(2, 1, _text(subtitle)).alignment = Alignment(horizontal='left')
    for index, header in enumerate(headers, 1):
        cell = sheet.cell(4, index, _text(header))
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='4F81BD')
        cell.alignment = Alignment(horizontal='center', vertical='center')
    for row_index, row in enumerate(rows, 5):
        for column_index, header in enumerate(headers, 1):
            cell = sheet.cell(row_index, column_index, _text(row.get(header, '')))
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    sheet.freeze_panes = 'A5'
    sheet.auto_filter.ref = f'A4:{get_column_letter(max(1, len(headers)))}{max(4, len(rows) + 4)}'
    for index, header in enumerate(headers, 1):
        widest = max([len(str(header))] + [len(str(_text(row.get(header, '')))) for row in rows])
        sheet.column_dimensions[get_column_letter(index)].width = min(max(widest + 2, 12), 35)
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def _docx(title, subtitle, headers, rows):
    document = Document()
    document.core_properties.title = title
    heading = document.add_heading(title, 0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(subtitle)
    if not rows:
        document.add_paragraph('当前筛选条件下没有可导出的演示数据。')
    else:
        table = document.add_table(rows=1, cols=len(headers))
        table.style = 'Table Grid'
        for cell, header in zip(table.rows[0].cells, headers):
            cell.text = str(_text(header))
        for row in rows:
            cells = table.add_row().cells
            for cell, header in zip(cells, headers):
                cell.text = str(_text(row.get(header, '')))
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def _payments(context):
    filters, start, end = context['filters'], *_date_range(context['filters'])
    records = []
    keyword = str(filters.get('keyword') or '').casefold()
    for payment in context['rows']['payments']:
        contract = context['contracts'].get(payment.get('contract_id'))
        if not contract or not _match_contract(contract, context, apply_dates=False):
            continue
        payment_date = payment.get('payment_date')
        if (start or end) and not _in_range(payment_date, start, end):
            continue
        if filters.get('payment_type') and payment.get('payment_type') != filters['payment_type']:
            continue
        if filters.get('currency') and payment.get('currency') != filters['currency']:
            continue
        if filters.get('export_type') and contract.get('export_type') != filters['export_type']:
            continue
        customer = context['customers'].get(contract.get('customer_id'), {})
        searchable = ' '.join(str(part or '') for part in (customer.get('name'), contract.get('contract_no'), payment.get('payment_type'))).casefold()
        if keyword and keyword not in searchable:
            continue
        records.append({'记录类型': '合同分摊', '客户': customer.get('name'), '合同号': contract.get('contract_no'),
                        '回款日期': payment_date, '回款类型': payment.get('payment_type'), '金额': payment.get('amount'),
                        '币种': payment.get('currency'), '人民币金额': payment.get('amount_rmb')})
    return records


def _rate_lookup(context, currency, when, latest=False):
    if currency == 'RMB':
        return 1
    month = str(when or '')[:7]
    if not month:
        return None
    choices = [row for row in context['rows']['exchange_rates']
               if row.get('currency_pair') == 'USD/CNY' and row.get('effective_month')]
    own = {row['effective_month']: row.get('rate') for row in choices if row.get('owner_id') == context['user']}
    shared = {row['effective_month']: row.get('rate') for row in choices if row.get('owner_id') == 'demo-shared'}
    rates = {**shared, **own}
    if latest:
        available = [key for key in rates if key <= month]
        return rates[max(available)] if available else None
    return rates.get(month)


def _component_rmb(context, currency, components):
    """Return None when any original-currency component lacks its required rate."""
    total = 0
    for amount, occurred_on, stored_rmb in components:
        if currency == 'RMB':
            total += amount
        elif stored_rmb is not None:
            total += stored_rmb
        else:
            rate = _rate_lookup(context, currency, occurred_on)
            if rate is None:
                return None
            total += amount * rate
    return total


def _customer_monthly(context):
    as_of = context['filters'].get('as_of_date') or date.today().isoformat()
    iso_date(as_of, '截至日期')
    month_start = as_of[:8] + '01'
    buckets = defaultdict(lambda: {
        'orders': [], 'monthly_payments': [], 'shipments': [], 'payments': [],
    })

    def bucket(customer_id, currency):
        return buckets[(customer_id, currency)]

    for contract in context['rows']['contracts']:
        if not _match_contract(contract, context, apply_dates=False):
            continue
        signed_on = contract.get('contract_date')
        if not signed_on or not (month_start <= signed_on <= as_of):
            continue
        for sheet in context['rows']['contact_sheets']:
            if sheet.get('contract_id') == contract['id'] and _match_sheet(sheet, context):
                bucket(contract['customer_id'], contract['currency'])['orders'].append(
                    ((sheet.get('quantity') or 0) * (sheet.get('unit_price') or 0), signed_on, None)
                )
    for shipment in context['rows']['shipments']:
        if not _match_shipment(shipment, context) or shipment.get('status') != '已发货':
            continue
        shipped_on = shipment.get('shipment_date')
        if not shipped_on or shipped_on > as_of:
            continue
        contract = context['contracts'][shipment['contract_id']]
        for item in context['rows']['shipment_items']:
            if item.get('shipment_id') == shipment['id'] and _match_sheet(context['sheets'].get(item.get('contact_sheet_id'), {}), context):
                amount = item.get('amount') or (item.get('shipped_quantity') or 0) * (item.get('unit_price') or 0)
                bucket(contract['customer_id'], contract['currency'])['shipments'].append((amount, shipped_on, None))
    receipt_ids = set()
    for receipt in context['rows']['payment_receipts']:
        paid_on, customer_id, currency = receipt.get('payment_date'), receipt.get('customer_id'), receipt.get('currency')
        if not paid_on or paid_on > as_of or not customer_id or not currency:
            continue
        receipt_ids.add(receipt['id'])
        target = bucket(customer_id, currency)
        component = (receipt.get('total_amount') or 0, paid_on, receipt.get('amount_rmb'))
        target['payments'].append(component)
        if paid_on >= month_start:
            target['monthly_payments'].append(component)
    # A payment with a receipt header is an allocation, not another receipt.
    for payment in context['rows']['payments']:
        if payment.get('receipt_id') in receipt_ids or payment.get('receipt_id'):
            continue
        contract = context['contracts'].get(payment.get('contract_id'))
        paid_on = payment.get('payment_date')
        if not contract or not _match_contract(contract, context, apply_dates=False) or not paid_on or paid_on > as_of:
            continue
        target = bucket(contract['customer_id'], contract['currency'])
        component = (payment.get('amount') or 0, paid_on, payment.get('amount_rmb'))
        target['payments'].append(component)
        if paid_on >= month_start:
            target['monthly_payments'].append(component)

    rows = []
    for (customer_id, currency), values in sorted(buckets.items(), key=lambda item: (context['customers'].get(item[0][0], {}).get('name', ''), item[0][1])):
        original = {key: sum(component[0] for component in values[key]) for key in values}
        receivable = max(0, original['shipments'] - original['payments'])
        prepayment = max(0, original['payments'] - original['shipments'])
        if not any((original['orders'], original['monthly_payments'], receivable, prepayment)):
            continue
        latest_rate = _rate_lookup(context, currency, as_of, latest=True)
        rows.append({
            '客户': context['customers'].get(customer_id, {}).get('name'), '币种': currency,
            '本月订单原币': original['orders'], '本月回款原币': original['monthly_payments'],
            '累计已发货原币': original['shipments'], '累计已回款原币': original['payments'],
            '当前应收原币': receivable, '客户预存款原币': prepayment, '期末汇率': latest_rate,
            '本月订单人民币': _component_rmb(context, currency, values['orders']),
            '本月回款人民币': _component_rmb(context, currency, values['monthly_payments']),
            '累计已发货人民币': _component_rmb(context, currency, values['shipments']),
            '累计已回款人民币': _component_rmb(context, currency, values['payments']),
            '当前应收人民币': None if latest_rate is None else receivable * latest_rate,
            '客户预存款人民币': None if latest_rate is None else prepayment * latest_rate,
            '汇率状态': '正常' if latest_rate is not None else '缺少汇率',
        })
    return rows


def _rmb_invoice_rows(context):
    filters = context['filters']
    scoped_shipment = _ids(filters, 'shipment_id') is not None or _ids(filters, 'shipment_group_id') is not None
    links_by_shipment = defaultdict(list)
    invoices = {row['id']: row for row in context['rows']['invoices']}
    for link in context['rows']['invoice_shipments']:
        if link.get('invoice_id') in invoices:
            links_by_shipment[link.get('shipment_id')].append(link)
    rows = []
    if scoped_shipment:
        for shipment in context['rows']['shipments']:
            if not _match_shipment(shipment, context):
                continue
            contract = context['contracts'][shipment['contract_id']]
            if contract.get('currency') != 'RMB' and contract.get('export_type') != '转口':
                continue
            customer = context['customers'].get(contract.get('customer_id'), {})
            items = [item for item in context['rows']['shipment_items'] if item.get('shipment_id') == shipment['id']]
            shipment_total = sum(item.get('amount') or (item.get('shipped_quantity') or 0) * (item.get('unit_price') or 0) for item in items)
            for item in items:
                sheet = context['sheets'].get(item.get('contact_sheet_id'), {})
                if not _match_sheet(sheet, context):
                    continue
                batch = context['batches'].get(item.get('batch_id'), {})
                amount = item.get('amount') or (item.get('shipped_quantity') or 0) * (item.get('unit_price') or 0)
                link = links_by_shipment.get(shipment['id'], [None])[0]
                invoice = invoices.get(link.get('invoice_id')) if link else None
                allocated = (link.get('allocated_amount') * amount / shipment_total) if link and shipment_total else None
                rows.append({'合同号': contract.get('contract_no'), '客户': customer.get('name'), '联系单号': sheet.get('contact_sheet_no'),
                             '物料号': sheet.get('material_no'), '产品名称': sheet.get('product_name'), '规格': sheet.get('specification'),
                             '数量': item.get('shipped_quantity'), '单位': sheet.get('unit'), '单价': item.get('unit_price'), '金额': amount,
                             '发货单号': shipment.get('shipment_no'), '发货日期': shipment.get('shipment_date'), '批号': batch.get('batch_no'),
                             '发票号': invoice.get('invoice_no') if invoice else None, '发票状态': invoice.get('status') if invoice else '未申请',
                             '分摊开票金额': allocated, '备注': invoice.get('notes') if invoice else sheet.get('notes')})
        return rows
    for sheet in context['rows']['contact_sheets']:
        contract = context['contracts'].get(sheet.get('contract_id'))
        if not contract or not _match_contract(contract, context) or not _match_sheet(sheet, context):
            continue
        if contract.get('currency') != 'RMB' and contract.get('export_type') != '转口':
            continue
        customer = context['customers'].get(contract.get('customer_id'), {})
        rows.append({'合同号': contract.get('contract_no'), '客户': customer.get('name'), '联系单号': sheet.get('contact_sheet_no'),
                     '物料号': sheet.get('material_no'), '产品名称': sheet.get('product_name'), '规格': sheet.get('specification'),
                     '数量': sheet.get('quantity'), '单位': sheet.get('unit'), '单价': sheet.get('unit_price'),
                     '金额': (sheet.get('quantity') or 0) * (sheet.get('unit_price') or 0), '发货单号': None,
                     '发货日期': None, '批号': None, '发票号': None, '发票状态': '未申请', '分摊开票金额': None, '备注': sheet.get('notes')})
    return rows


def export_file(store, command, filters, user='demo-sales-1'):
    """Return ``(filename, MIME type, content)`` for one supported local export."""
    if command not in COMMANDS:
        raise ApiError('未知导出命令', 404, 'unknown_export')
    context = _context(store, {} if filters is None else filters, user)
    subtitle = '本文件仅含虚构演示数据，由本地 SQLite 数据库生成。'
    if command == 'export-production':
        title, rows = '生产协调与催包材统计表', _operational_rows(context, 'production')
        headers = ['联系单号', '合同号', '客户', '产品名称', '规格', '批号汇总', '生产/排产日期', '有效期', 'QA审批日期',
                   '实际入库日期', '包装/包材', '数量', '单位', '已发货数量', '交付说明']
    elif command == 'export-monthly':
        title, rows = '订单月度汇总表', _line_rows(context)
    elif command == 'export-original-contracts':
        title = '正本合同登记统计表'
        rows = []
        for contract in context['rows']['contracts']:
            if _match_contract(contract, context):
                customer = context['customers'].get(contract.get('customer_id'), {})
                rows.append({'客户': customer.get('name'), '合同号': contract.get('contract_no'), '签署日期': contract.get('contract_date'),
                             '币种': contract.get('currency'), '贸易术语': contract.get('incoterm'), '贸易方式': contract.get('export_type'), '状态': contract.get('status')})
    elif command == 'export-helper':
        title, rows = '下联系单辅助表', _operational_rows(context, 'helper')
        headers = ['联系单号', '产品名称', '规格', '批号汇总', '生产日期', '有效期', '数量', '单位', '包装/包材',
                   '参考联系单', '质量标准', '批次数量说明']
    elif command == 'export-payments':
        title, rows = '回款明细表', _payments(context)
    elif command == 'export-customer-monthly-sales':
        title, rows = '客户实际销售、回款及应收汇总', _customer_monthly(context)
        headers = ['客户', '币种', '本月订单原币', '本月回款原币', '累计已发货原币', '累计已回款原币', '当前应收原币',
                   '客户预存款原币', '期末汇率', '本月订单人民币', '本月回款人民币', '累计已发货人民币', '累计已回款人民币',
                   '当前应收人民币', '客户预存款人民币', '汇率状态']
    elif command == 'export-rmb-invoice':
        title, rows = '人民币增值税开票申请表', _rmb_invoice_rows(context)
        headers = ['合同号', '客户', '联系单号', '物料号', '产品名称', '规格', '数量', '单位', '单价', '金额', '发货单号',
                   '发货日期', '批号', '发票号', '发票状态', '分摊开票金额', '备注']
    elif command == 'export-shipment-plan':
        title, rows = '发货安排表', _line_rows(context, shipments_only=True)
        headers = ['客户', '合同号', '联系单号', '产品名称', '批号', '发货单号', '发货日期', '发货状态', '数量', '单位', '单价', '金额', '币种']
    elif command == 'export-receipt-confirmation':
        title, rows = '收货确认函（演示）', _line_rows(context, shipments_only=True)
        headers = ['客户', '合同号', '联系单号', '产品名称', '批号', '发货单号', '发货日期', '数量', '单位', '币种']
        return '收货确认函-演示.docx', DOCX_MIME, _docx(title, subtitle, headers, rows)
    else:
        title, rows = '全部发货明细', []
        selected_shipments = {shipment['id'] for shipment in context['rows']['shipments'] if _match_shipment(shipment, context)}
        selected_sheets = _ids(context['filters'], 'contact_sheet_ids')
        for profit in store.query('v_shipment_profit', user=user):
            if profit.get('shipment_id') not in selected_shipments:
                continue
            sheet = next((row for row in context['rows']['contact_sheets'] if row.get('contact_sheet_no') == profit.get('contact_sheet_no')
                          and (selected_sheets is None or row.get('id') in selected_sheets)), None)
            if selected_sheets is not None and sheet is None:
                continue
            rows.append({'发货单号': profit.get('shipment_no'), '发货日期': profit.get('shipment_date'), '合同号': profit.get('contract_no'),
                         '客户': profit.get('customer_name'), '联系单号': profit.get('contact_sheet_no'), '批号': profit.get('batch_no'),
                         '物料号': profit.get('material_no'), '产品名称': profit.get('product_name'), '规格': profit.get('specification'),
                         '数量': profit.get('shipped_quantity'), '单位': profit.get('unit'), '单价': profit.get('unit_price'),
                         '金额': profit.get('sales_amount'), '币种': profit.get('currency'), '汇率': profit.get('exchange_rate'),
                         '利润（人民币）': profit.get('profit'), '利润状态': '待计算' if profit.get('profit') is None else '已计算'})
        headers = ['发货单号', '发货日期', '合同号', '客户', '联系单号', '批号', '物料号', '产品名称', '规格', '数量', '单位', '单价',
                   '金额', '币种', '汇率', '利润（人民币）', '利润状态']
    if 'headers' not in locals():
        headers = list(rows[0]) if rows else ['客户', '合同号', '联系单号', '产品名称']
    return f'{title}-演示.xlsx', XLSX_MIME, _xlsx(title, subtitle, headers, rows)
